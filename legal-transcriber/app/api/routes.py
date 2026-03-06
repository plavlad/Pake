from __future__ import annotations

import asyncio
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import APIRouter, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from app.core.config import UPLOAD_DIR, settings
from app.services.job_store import job_store
from app.services.transcription import compose_formatted_transcript, extract_speakers, pipeline


router = APIRouter(prefix="/api", tags=["transcription"])
task_semaphore = asyncio.Semaphore(settings.task_concurrency)


class SpeakerMapPayload(BaseModel):
    speaker_map: dict[str, str] = Field(default_factory=dict)


def _sanitize_filename(filename: str) -> str:
    safe = "".join(char for char in filename if char.isalnum() or char in {"-", "_", ".", " "}).strip()
    return safe or "audio"


def _build_public_task(task: dict[str, Any]) -> dict[str, Any]:
    public_task = {
        "id": task["id"],
        "source_filename": task["source_filename"],
        "status": task["status"],
        "stage": task["stage"],
        "progress": task["progress"],
        "error": task["error"],
        "created_at": task["created_at"],
        "updated_at": task["updated_at"],
        "result": None,
    }

    result = task.get("result")
    if not result:
        return public_task

    speaker_map = task.get("speaker_map") or result.get("speaker_map") or {}
    segments: list[dict[str, Any]] = []
    for segment in result.get("segments", []):
        item = dict(segment)
        item["display_speaker"] = speaker_map.get(item["speaker"], item["speaker"])
        segments.append(item)

    public_result = dict(result)
    public_result["speaker_map"] = speaker_map
    public_result["segments"] = segments
    public_result["speakers"] = extract_speakers(result.get("segments", []), speaker_map)
    public_result["formatted_text"] = compose_formatted_transcript(result.get("segments", []), speaker_map)

    public_task["result"] = public_result
    return public_task


async def _save_upload(file: UploadFile) -> Path:
    suffix = Path(file.filename or "").suffix.lower()
    safe_name = _sanitize_filename(Path(file.filename or "audio").stem)
    target_path = UPLOAD_DIR / f"{safe_name}_{uuid.uuid4().hex}{suffix}"

    size_limit = settings.max_upload_size_mb * 1024 * 1024
    total_size = 0

    with target_path.open("wb") as target:
        while chunk := await file.read(1024 * 1024):
            total_size += len(chunk)
            if total_size > size_limit:
                target_path.unlink(missing_ok=True)
                raise HTTPException(
                    status_code=413,
                    detail=f"Файл слишком большой. Лимит: {settings.max_upload_size_mb} MB.",
                )
            target.write(chunk)

    await file.close()
    return target_path


async def _process_task(task_id: str) -> None:
    async with task_semaphore:
        task = job_store.get(task_id)
        if task is None:
            return

        def report(stage: str, progress: int) -> None:
            job_store.mark_processing(task_id, stage=stage, progress=progress)

        try:
            report("Подготовка", 3)
            result = await asyncio.to_thread(pipeline.run, task["stored_path"], report)
            result["speaker_map"] = task.get("speaker_map") or result.get("speaker_map") or {}
            job_store.mark_completed(task_id, result)
        except Exception as exc:
            current = job_store.get(task_id) or {}
            job_store.mark_failed(task_id, error=str(exc), stage=current.get("stage"))


def _build_export_text(task: dict[str, Any]) -> str:
    result = _build_public_task(task).get("result") or {}
    header = f"Файл: {task['source_filename']}\nСтатус: готово\n"
    return header + "\n" + result.get("formatted_text", "")


def _build_docx_bytes(task: dict[str, Any]) -> bytes:
    document = Document()
    document.add_heading("Протокол аудиозаписи", level=1)
    document.add_paragraph(f"Исходный файл: {task['source_filename']}")

    result = _build_public_task(task).get("result") or {}
    for line in result.get("formatted_text", "").splitlines():
        document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.read()


@router.get("/health")
async def health() -> dict[str, str]:
    return {"status": "ok"}


@router.post("/tasks")
async def create_task(file: UploadFile = File(...)) -> dict[str, Any]:
    if not file.filename:
        raise HTTPException(status_code=400, detail="Не удалось определить имя файла.")

    stored_path = await _save_upload(file)
    task = job_store.create_task(filename=file.filename, stored_path=stored_path)
    asyncio.create_task(_process_task(task["id"]))
    return _build_public_task(task)


@router.get("/tasks/{task_id}")
async def get_task(task_id: str) -> dict[str, Any]:
    task = job_store.get(task_id)
    if task is None:
        raise HTTPException(status_code=404, detail="Задача не найдена.")
    return _build_public_task(task)


@router.put("/tasks/{task_id}/speaker-map")
async def update_speaker_map(task_id: str, payload: SpeakerMapPayload) -> dict[str, Any]:
    task = job_store.get(task_id)
    if task is None:
        raise HTTPException(status_code=404, detail="Задача не найдена.")
    if not task.get("result"):
        raise HTTPException(status_code=409, detail="Результат еще не готов.")

    normalized_map = {key: value.strip() or key for key, value in payload.speaker_map.items()}
    updated = job_store.update_speaker_map(task_id, normalized_map)
    return _build_public_task(updated)


@router.get("/tasks/{task_id}/export")
async def export_task(
    task_id: str,
    format: str = Query(default="txt", pattern="^(txt|docx)$"),
) -> StreamingResponse:
    task = job_store.get(task_id)
    if task is None:
        raise HTTPException(status_code=404, detail="Задача не найдена.")
    if task.get("status") != "completed":
        raise HTTPException(status_code=409, detail="Экспорт доступен только после завершения обработки.")

    base_name = Path(task["source_filename"]).stem

    if format == "txt":
        content = _build_export_text(task).encode("utf-8")
        return StreamingResponse(
            BytesIO(content),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{base_name}.txt"'},
        )

    content = _build_docx_bytes(task)
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{base_name}.docx"'},
    )
