from __future__ import annotations

import asyncio
import uuid
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles

from .config import (
    EXPORT_DIR,
    MAX_UPLOAD_SIZE_BYTES,
    STATIC_DIR,
    SUPPORTED_AUDIO_EXTENSIONS,
    UPLOAD_DIR,
)
from .models import ExportRequest, JobResult, JobState
from .pipeline import LegalTranscriptionPipeline, PipelineError
from .utils import render_plain_text

app = FastAPI(title="Legal Local Transcriber", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

STATIC_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

jobs: dict[str, JobState] = {}
jobs_lock = asyncio.Lock()
processing_semaphore = asyncio.Semaphore(1)
pipeline = LegalTranscriptionPipeline()


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


async def _update_job(
    job_id: str,
    *,
    stage: str | None = None,
    progress: int | None = None,
    message: str | None = None,
    error: str | None = None,
    result: JobResult | None = None,
) -> None:
    async with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            return
        if stage is not None:
            job.stage = stage  # type: ignore[assignment]
        if progress is not None:
            job.progress = max(0, min(100, progress))
        if message is not None:
            job.message = message
        if error is not None:
            job.error = error
        if result is not None:
            job.result = result
        job.updated_at = _utc_now_iso()
        jobs[job_id] = job


async def _run_job(job_id: str, audio_path: Path) -> None:
    loop = asyncio.get_running_loop()

    # Перенаправляем update через корректный event loop.
    def on_stage_threadsafe(stage: str, progress: int, message: str) -> None:
        asyncio.run_coroutine_threadsafe(
            _update_job(job_id, stage=stage, progress=progress, message=message),
            loop,
        )

    try:
        await _update_job(
            job_id,
            stage="queued",
            progress=1,
            message="Ожидание свободного AI-воркера",
        )
        async with processing_semaphore:
            raw_result = await asyncio.to_thread(
                pipeline.run,
                str(audio_path),
                on_stage_threadsafe,
            )
            result = JobResult.model_validate(raw_result)
            await _update_job(
                job_id,
                stage="completed",
                progress=100,
                message="Обработка завершена",
                result=result,
            )
    except PipelineError as exc:
        await _update_job(
            job_id,
            stage="failed",
            progress=100,
            message="Ошибка обработки",
            error=str(exc),
        )
    except Exception as exc:  # noqa: BLE001
        await _update_job(
            job_id,
            stage="failed",
            progress=100,
            message="Критическая ошибка",
            error=f"Неожиданная ошибка: {exc}",
        )


@app.get("/")
async def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/health")
async def health() -> dict:
    return {"status": "ok"}


@app.post("/api/upload")
async def upload_audio(file: UploadFile = File(...)) -> dict:
    filename = file.filename or "audio"
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_AUDIO_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Неподдерживаемый формат файла.")

    job_id = uuid.uuid4().hex
    safe_filename = f"{job_id}{suffix}"
    save_path = UPLOAD_DIR / safe_filename
    total_size = 0
    with save_path.open("wb") as output:
        while True:
            chunk = await file.read(1024 * 1024)
            if not chunk:
                break
            total_size += len(chunk)
            if total_size > MAX_UPLOAD_SIZE_BYTES:
                output.close()
                save_path.unlink(missing_ok=True)
                raise HTTPException(
                    status_code=413,
                    detail=(
                        f"Файл слишком большой. Максимум: "
                        f"{MAX_UPLOAD_SIZE_BYTES // (1024 * 1024)} МБ."
                    ),
                )
            output.write(chunk)

    job = JobState(
        job_id=job_id,
        filename=filename,
        stage="queued",
        progress=0,
        message="Файл загружен, задача в очереди",
    )
    async with jobs_lock:
        jobs[job_id] = job

    asyncio.create_task(_run_job(job_id, save_path))
    return {"job_id": job_id}


@app.get("/api/status/{job_id}")
async def get_status(job_id: str) -> dict:
    async with jobs_lock:
        job = jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Задача не найдена.")
    return job.model_dump()


@app.get("/api/result/{job_id}")
async def get_result(job_id: str) -> dict:
    async with jobs_lock:
        job = jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Задача не найдена.")
    if job.stage != "completed" or not job.result:
        raise HTTPException(status_code=409, detail="Результат пока не готов.")
    return job.result.model_dump()


@app.post("/api/export/{job_id}")
async def export_result(job_id: str, request: ExportRequest, format: str = "txt") -> FileResponse:
    async with jobs_lock:
        job = jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Задача не найдена.")
    if job.stage != "completed" or not job.result:
        raise HTTPException(status_code=409, detail="Результат пока не готов.")

    segments = job.result.segments
    text = render_plain_text(segments, speaker_map=request.speaker_map)
    export_id = uuid.uuid4().hex

    if format == "txt":
        out_path = EXPORT_DIR / f"{job_id}_{export_id}.txt"
        out_path.write_text(text, encoding="utf-8")
        return FileResponse(
            out_path,
            media_type="text/plain; charset=utf-8",
            filename=f"transcript_{job_id}.txt",
        )

    if format == "docx":
        out_path = EXPORT_DIR / f"{job_id}_{export_id}.docx"
        doc = Document()
        doc.add_heading("Протокол расшифровки", level=1)
        for line in text.splitlines():
            doc.add_paragraph(line)
        doc.save(out_path)
        return FileResponse(
            out_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=f"transcript_{job_id}.docx",
        )

    raise HTTPException(status_code=400, detail="Поддерживаются только txt и docx.")
