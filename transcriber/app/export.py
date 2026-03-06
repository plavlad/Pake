"""Экспорт результатов в .txt и .docx."""

from __future__ import annotations

import io
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def export_txt(text: str, speaker_map: dict[str, str] | None = None) -> str:
    """Возвращает текст с заменёнными именами спикеров."""
    return _apply_speaker_map(text, speaker_map)


def export_docx(
    text: str,
    speaker_map: dict[str, str] | None = None,
    title: str = "Протокол",
) -> bytes:
    """Генерирует .docx файл из текста транскрипции."""
    text = _apply_speaker_map(text, speaker_map)

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)

    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")

    lines = text.strip().split("\n")
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph("")
            continue

        p = doc.add_paragraph()

        ts_match = re.match(r"(\[[\d:]+\s*-\s*[\d:]+\])\s*(.*)", line)
        if ts_match:
            timecode = ts_match.group(1)
            rest = ts_match.group(2)

            run_tc = p.add_run(timecode + " ")
            run_tc.font.size = Pt(10)
            run_tc.font.color.rgb = RGBColor(128, 128, 128)

            speaker_match = re.match(r"([^:]+):\s*(.*)", rest)
            if speaker_match:
                speaker_name = speaker_match.group(1)
                speech = speaker_match.group(2)

                run_spk = p.add_run(speaker_name + ": ")
                run_spk.bold = True
                run_spk.font.size = Pt(12)

                run_text = p.add_run(speech)
                run_text.font.size = Pt(12)
            else:
                run_text = p.add_run(rest)
                run_text.font.size = Pt(12)
        else:
            run_text = p.add_run(line)
            run_text.font.size = Pt(12)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def _apply_speaker_map(text: str, speaker_map: dict[str, str] | None) -> str:
    """Заменяет метки спикеров на пользовательские имена."""
    if not speaker_map:
        return text
    for original, replacement in speaker_map.items():
        if original and replacement and original != replacement:
            text = text.replace(original, replacement)
    return text
